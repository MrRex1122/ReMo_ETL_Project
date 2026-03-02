import os
import tempfile
import unittest
from codecs import BOM_UTF8
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from snapshot_export import (
    ARTICLE_DUPLICATE_COLUMN,
    NAME_DUPLICATE_COLUMN,
    build_duplicate_report_from_csv,
    build_xlsx_from_csv_streaming,
    get_snapshot_xlsx_status,
    stage_public_export,
)


class SnapshotExportTests(unittest.TestCase):
    def test_build_duplicate_report_from_csv_matches_expected_stats(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_csv = root / "catalog.csv"
            duplicate_csv = root / "duplicates.csv"
            pd.DataFrame(
                [
                    {"Наименование": "Кабель UTP cat6", "Артикул": "A-1", "Цена розничная": 10},
                    {"Наименование": "Кабель UTP cat6", "Артикул": "", "Цена розничная": 11},
                    {"Наименование": "Автомат 16A", "Артикул": "A-1", "Цена розничная": 12},
                ]
            ).to_csv(source_csv, sep=";", encoding="utf-8", index=False)

            stats, resolved_duplicate_csv = build_duplicate_report_from_csv(
                source_csv,
                duplicate_csv_path=duplicate_csv,
                chunksize=2,
            )

            self.assertEqual(stats["rows_total"], 3)
            self.assertEqual(stats["duplicates_total"], 3)
            self.assertEqual(stats["duplicates_by_article"], 2)
            self.assertEqual(stats["duplicates_by_name"], 2)
            self.assertEqual(resolved_duplicate_csv, duplicate_csv)
            out_df = pd.read_csv(duplicate_csv, sep=";", encoding="utf-8-sig")
            self.assertEqual(len(out_df), 3)
            self.assertIn(ARTICLE_DUPLICATE_COLUMN, out_df.columns)
            self.assertIn(NAME_DUPLICATE_COLUMN, out_df.columns)

    def test_stage_public_export_streams_and_adds_utf8_bom(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_csv = root / "catalog.csv"
            public_dir = root / "public"
            source_csv.write_text("Наименование;Артикул\nКабель;A-1\n", encoding="utf-8")

            public_path, public_url = stage_public_export(
                source_csv,
                "catalog_snapshot.csv",
                public_dir=public_dir,
                add_utf8_bom=True,
            )

            self.assertTrue(public_path.exists())
            self.assertEqual(public_url, "/app/static/exports/catalog_snapshot.csv")
            self.assertTrue(public_path.read_bytes().startswith(BOM_UTF8))

    def test_build_xlsx_from_csv_streaming_creates_valid_workbook(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_csv = root / "catalog.csv"
            target_xlsx = root / "catalog.xlsx"
            source_csv.write_text(
                "Наименование;Артикул;Цена розничная\nКабель UTP cat6;A-1;100\n",
                encoding="utf-8",
            )

            build_xlsx_from_csv_streaming(source_csv, target_xlsx)

            self.assertTrue(target_xlsx.exists())
            workbook = load_workbook(target_xlsx, read_only=True)
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            workbook.close()
            self.assertEqual(rows[0], ("Наименование", "Артикул", "Цена розничная"))
            self.assertEqual(rows[1], ("Кабель UTP cat6", "A-1", "100"))

    def test_get_snapshot_xlsx_status_reports_stale_partial(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            root = Path(tmp_dir)
            source_csv = root / "catalog.csv"
            target_xlsx = root / "catalog.xlsx"
            part_path = Path(f"{target_xlsx}.part")
            source_csv.write_text("Наименование;Артикул\nКабель;A-1\n", encoding="utf-8")
            part_path.write_text("building", encoding="utf-8")
            stale_time = part_path.stat().st_mtime - 3600
            os.utime(part_path, (stale_time, stale_time))

            status, started_at = get_snapshot_xlsx_status(
                source_csv,
                target_xlsx,
                stale_after_seconds=10,
            )

            self.assertEqual(status, "failed_stale")
            self.assertIsNotNone(started_at)


if __name__ == "__main__":
    unittest.main()
